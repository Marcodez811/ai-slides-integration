"""Editable, source-attributed PPTX materialization for validated layouts."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
import shutil
import subprocess

from .layout import LayoutValidationError, validate_layout
from .models import PhysicalSlideLayout, PresentationLayout, RenderDiagnostic, RenderReport


class PresentationRenderError(RuntimeError):
    """Raised when a layout cannot be materialized as a valid PPTX package."""


def render_presentation(layout: PresentationLayout | Sequence[PhysicalSlideLayout], output_path: str | Path) -> RenderReport:
    """Render an editable 16:9 deck and verify that the package reopens."""
    presentation_layout = layout if isinstance(layout, PresentationLayout) else PresentationLayout(slides=list(layout))
    diagnostics = [diagnostic for slide in presentation_layout.slides for diagnostic in validate_layout(slide)]
    theme = presentation_layout.slides[0].theme
    if not _font_is_locally_available(theme.title_font):
        diagnostics.append(RenderDiagnostic(
            severity="warning",
            code="FONT_NOT_INSTALLED_LOCALLY",
            message=(
                f"Requested font {theme.title_font!r} is not installed locally; "
                f"PowerPoint may use {theme.fallback_font!r}."
            ),
        ))
    errors = [diagnostic for diagnostic in diagnostics if diagnostic.severity == "error"]
    if errors:
        raise LayoutValidationError("; ".join(diagnostic.message for diagnostic in errors))
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches, Pt
    except ImportError as error:  # pragma: no cover - optional package boundary
        raise PresentationRenderError("python-pptx is required to render PPTX files") from error

    target = Path(output_path)
    if target.suffix.lower() != ".pptx":
        raise PresentationRenderError("output_path must end with .pptx")
    target.parent.mkdir(parents=True, exist_ok=True)
    first_theme = presentation_layout.slides[0].theme
    deck = Presentation()
    deck.slide_width = Inches(first_theme.slide_width_inches)
    deck.slide_height = Inches(first_theme.slide_height_inches)
    blank_layout = deck.slide_layouts[6]
    for slide_layout in presentation_layout.slides:
        slide = deck.slides.add_slide(blank_layout)
        _set_background(slide, slide_layout.theme.background_color, RGBColor)
        for element in slide_layout.elements:
            _render_element(slide, element, slide_layout.theme, Inches, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR, CategoryChartData, XL_CHART_TYPE, XL_LEGEND_POSITION, OxmlElement, diagnostics, slide_layout)
    deck.save(target)
    try:
        reopened = Presentation(target)
        verified = len(reopened.slides) == len(presentation_layout.slides)
    except Exception as error:  # pragma: no cover - defensive verification boundary
        diagnostics.append(RenderDiagnostic(severity="error", code="package_reopen", message=f"PPTX package could not be reopened: {error}"))
        verified = False
    if not verified:
        raise PresentationRenderError("PPTX package reopen verification failed")
    semantic_ids = {slide.semantic_slide_id or slide.slide_id for slide in presentation_layout.slides}
    source_assets = sorted({element.payload.path for slide in presentation_layout.slides for element in slide.elements if element.payload and element.payload.kind == "image" and element.payload.element_index >= 0 and element.payload.path})
    generated_assets = sorted({element.payload.path for slide in presentation_layout.slides for element in slide.elements if element.payload and element.payload.kind == "image" and element.payload.element_index < 0 and element.payload.path})
    return RenderReport(output_path=str(target), slides_rendered=len(presentation_layout.slides), semantic_slides_rendered=len(semantic_ids), physical_slide_ids=[slide.slide_id for slide in presentation_layout.slides], source_footer_count=sum(bool(slide.source_attributions) for slide in presentation_layout.slides), diagnostics=diagnostics, verified=True, artifact_paths={"pptx": str(target)}, source_assets=source_assets, generated_assets=generated_assets)


def _render_element(slide, element, theme, Inches, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR, CategoryChartData, XL_CHART_TYPE, XL_LEGEND_POSITION, OxmlElement, diagnostics, slide_layout) -> None:
    box = element.box
    x, y, width, height = (Inches(box.x), Inches(box.y), Inches(box.width), Inches(box.height))
    if element.kind == "image":
        assert element.payload is not None and element.payload.image_path is not None
        path = element.payload.image_path
        if not path.is_file():
            diagnostics.append(RenderDiagnostic(severity="warning", code="SOURCE_IMAGE_UNAVAILABLE", message="Image file was unavailable during rendering.", slide_id=slide_layout.semantic_slide_id, physical_slide_id=slide_layout.slide_id, element_index=element.element_index))
            _add_textbox(slide, "Source image unavailable.", element, theme, x, y, width, height, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR, OxmlElement)
            return
        _add_cropped_picture(slide, path, x, y, width, height)
        return
    if element.kind == "table":
        assert element.payload is not None and element.payload.table_payload is not None
        _add_table(slide, element, theme, x, y, width, height, Pt, RGBColor, MSO_ANCHOR, OxmlElement)
        return
    if element.kind == "chart":
        assert element.payload is not None
        _add_chart(slide, element, theme, x, y, width, height, Pt, RGBColor, CategoryChartData, XL_CHART_TYPE, XL_LEGEND_POSITION, OxmlElement)
        return
    text = element.text or (element.payload.text if element.payload else "")
    items = element.items or (element.payload.items if element.payload else None)
    _add_textbox(slide, text, element, theme, x, y, width, height, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR, OxmlElement, items=items)


def _add_textbox(slide, text, element, theme, x, y, width, height, Pt, RGBColor, PP_ALIGN, MSO_ANCHOR, OxmlElement, *, items=None) -> None:
    shape = slide.shapes.add_textbox(x, y, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    values = items if items is not None else [text]
    for index, value in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = value
        paragraph.alignment = PP_ALIGN.LEFT
        if items is not None:
            _set_real_bullet(paragraph, OxmlElement)
        font = paragraph.font
        font.name = theme.title_font if element.role in {"title", "section_title"} else theme.body_font
        font.bold = element.role in {"title", "section_title"}
        font.size = Pt(element.font_size_pt or theme.minimum_body_font_pt)
        font.color.rgb = RGBColor.from_string(theme.primary_color if element.role in {"title", "section_title"} else (theme.muted_color if element.role in {"caption", "footer"} else theme.text_color))
        _set_east_asian_font(paragraph, theme.east_asian_font, OxmlElement)


def _set_background(slide, color: str, RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


def _add_cropped_picture(slide, path: Path, x, y, width, height) -> None:
    """Fill a visual box while preserving source aspect ratio."""
    try:
        from PIL import Image
        with Image.open(path) as image:
            source_ratio = image.width / image.height
    except Exception:
        slide.shapes.add_picture(str(path), x, y, width=width, height=height)
        return
    target_ratio = width / height
    picture = slide.shapes.add_picture(str(path), x, y, width=width, height=height)
    if source_ratio > target_ratio:
        crop = (1 - target_ratio / source_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif source_ratio < target_ratio:
        crop = (1 - source_ratio / target_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop


def _add_table(slide, element, theme, x, y, width, height, Pt, RGBColor, MSO_ANCHOR, OxmlElement) -> None:
    assert element.payload is not None and element.payload.table_payload is not None
    payload = element.payload.table_payload
    table = slide.shapes.add_table(len(payload.rows), payload.column_count, x, y, width, height).table
    headers = set(payload.header_rows)
    # Write before merging: the exact source cell text remains editable in the
    # top-left PowerPoint cell of each source span.
    for row_index, row in enumerate(payload.rows):
        for column_index, source_cell in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = source_cell.text
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(theme.primary_color if row_index in headers else "FFFFFF")
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.name = theme.body_font
            paragraph.font.size = Pt(12)
            paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index in headers else theme.text_color)
            _set_east_asian_font(paragraph, theme.east_asian_font, OxmlElement)
    for row_index, row in enumerate(payload.rows):
        for column_index, source_cell in enumerate(row):
            if source_cell.row_span > 1 or source_cell.column_span > 1:
                last_row = min(len(payload.rows) - 1, row_index + source_cell.row_span - 1)
                last_column = min(payload.column_count - 1, column_index + source_cell.column_span - 1)
                table.cell(row_index, column_index).merge(table.cell(last_row, last_column))


def _add_chart(slide, element, theme, x, y, width, height, Pt, RGBColor, CategoryChartData, XL_CHART_TYPE, XL_LEGEND_POSITION, OxmlElement) -> None:
    assert element.payload is not None and element.payload.categories is not None and element.payload.series is not None and element.payload.chart_type is not None
    chart_data = CategoryChartData()
    chart_data.categories = element.payload.categories
    for series in element.payload.series:
        chart_data.add_series(series.name, series.values)
    chart_types = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA, "pie": XL_CHART_TYPE.PIE}
    chart = slide.shapes.add_chart(chart_types[element.payload.chart_type], x, y, width, height, chart_data).chart
    if element.payload.title:
        chart.has_title = True
        chart.chart_title.text_frame.text = element.payload.title
        title_paragraph = chart.chart_title.text_frame.paragraphs[0]
        title_paragraph.font.name = theme.body_font
        title_paragraph.font.size = Pt(12)
        _set_east_asian_font(title_paragraph, theme.east_asian_font, OxmlElement)
    chart.has_legend = len(element.payload.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    if element.payload.chart_type != "pie":
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.tick_labels.font.size = Pt(10)


def _set_real_bullet(paragraph, OxmlElement) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag.endswith("}buNone") or child.tag.endswith("}buChar") or child.tag.endswith("}buAutoNum"):
            p_pr.remove(child)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "•")
    p_pr.append(bullet)


def _set_east_asian_font(paragraph, typeface: str, OxmlElement) -> None:
    """Set both Latin and East Asian font XML so CJK text survives Office."""
    for run in paragraph.runs:
        run.font.name = run.font.name or typeface
        r_pr = run._r.get_or_add_rPr()
        east_asian = next((node for node in r_pr if node.tag.endswith("}ea")), None)
        if east_asian is None:
            east_asian = OxmlElement("a:ea")
            r_pr.append(east_asian)
        east_asian.set("typeface", typeface)


def _font_is_locally_available(font_name: str) -> bool:
    """Best-effort diagnostic only; unavailable fonts never block PPTX output."""
    executable = shutil.which("fc-match")
    if executable is None:
        return True
    try:
        result = subprocess.run(
            [executable, "--format=%{family}", font_name],
            check=False,
            capture_output=True,
            text=True,
            timeout=1,
        )
    except (OSError, subprocess.SubprocessError):
        return True
    families = {part.strip().casefold() for part in result.stdout.split(",") if part.strip()}
    return font_name.casefold() in families
