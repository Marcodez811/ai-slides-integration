"""Opt-in real-provider smoke test for validated semantic slide content."""

from __future__ import annotations

import asyncio
import os

import pytest
from docx import Document


_RUN_LIVE = os.environ.get("RUN_OPENAI_INTEGRATION_TESTS") == "1" and bool(
    os.environ.get("OPENAI_API_KEY")
)
pytestmark = pytest.mark.skipif(
    not _RUN_LIVE,
    reason="requires OPENAI_API_KEY and RUN_OPENAI_INTEGRATION_TESTS=1",
)


def test_openai_generates_validated_semantic_slide_content_from_tiny_docx(tmp_path) -> None:
    from presentation_pipeline import generate_plan
    from presentation_pipeline.artifacts import write_presentation_content_json
    from presentation_pipeline.planning import PresentationRequirements
    from presentation_pipeline.providers import OpenAIStructuredGenerator
    from presentation_pipeline.synthesis import (
        PresentationContent,
        build_presentation_content,
        build_slide_contexts,
        generate_slide_contents,
    )

    source_path = tmp_path / "briefing.docx"
    document = Document()
    document.add_heading("Quarterly update", level=1)
    document.add_paragraph("Revenue grew from 100 to 120 units in the second quarter.")
    document.add_paragraph("Continue the current growth plan.")
    document.save(source_path)

    requirements = PresentationRequirements(
        goal="Brief leadership on the quarterly update",
        audience="Leadership team",
        target_slide_count=2,
        presentation_type="briefing",
    )
    generator = OpenAIStructuredGenerator(model=os.environ.get("OPENAI_MODEL", "gpt-5.6-terra"))

    async def run() -> tuple[object, object, object]:
        plan = await generate_plan(
            [source_path], requirements, generator, extraction_workers=1, llm_concurrency=1
        )
        contexts = build_slide_contexts(plan)
        slides = await generate_slide_contents(contexts, generator, concurrency=1)
        return plan, contexts, build_presentation_content(slides, contexts)

    plan, contexts, content = asyncio.run(run())
    assert len(plan.outline.all_slides()) == requirements.target_slide_count
    assert len(contexts) == len(plan.outline.all_slides())
    assert [slide.slide_id for slide in content.slides] == [context.slide.slide_id for context in contexts]

    output_path = write_presentation_content_json(content, tmp_path / "slide-content.json")
    assert PresentationContent.model_validate_json(output_path.read_text(encoding="utf-8")) == content
