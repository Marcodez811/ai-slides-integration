"""Opt-in smoke test for a real OpenAI-backed DOCX-to-outline run.

This module is intentionally skipped unless both the credential and explicit
opt-in flag are available, so the ordinary test suite never sends network
requests or incurs model costs.
"""

from __future__ import annotations

import asyncio
import os

import pytest
from docx import Document


_RUN_LIVE = os.environ.get("RUN_OPENAI_INTEGRATION_TESTS") == "1" and bool(
    os.environ.get("OPENAI_API_KEY")
)
pytestmark = pytest.mark.skipif(
    not _RUN_LIVE,
    reason="requires OPENAI_API_KEY and RUN_OPENAI_INTEGRATION_TESTS=1",
)


def test_openai_generates_provenance_valid_json_outline_from_tiny_docx(tmp_path) -> None:
    from presentation_pipeline import generate_outline, write_outline_json
    from presentation_pipeline.corpus.batch import build_jobs, extract_batch
    from presentation_pipeline.indexing.builder import build_document_index
    from presentation_pipeline.planning import PresentationRequirements
    from presentation_pipeline.providers import OpenAIStructuredGenerator
    from presentation_pipeline.validation import CorpusLookup, validate_presentation_outline

    source_path = tmp_path / "briefing.docx"
    document = Document()
    document.add_heading("Quarterly update", level=1)
    document.add_paragraph("Revenue grew from 100 to 120 units in the second quarter.")
    document.add_paragraph("The recommendation is to continue the current growth plan.")
    document.save(source_path)

    requirements = PresentationRequirements(
        goal="Brief leadership on the quarterly update",
        audience="Leadership team",
        target_slide_count=2,
        presentation_type="briefing",
    )
    model = os.environ.get("OPENAI_MODEL", "gpt-5.6-terra")
    outline = asyncio.run(
        generate_outline(
            [source_path],
            requirements,
            OpenAIStructuredGenerator(model=model),
            extraction_workers=1,
            llm_concurrency=1,
        )
    )

    # Independently reconstruct the deterministic lookup to make the test's
    # provenance assertion explicit rather than relying only on the pipeline
    # contract.
    batch = extract_batch(build_jobs([source_path]), max_workers=1)
    assert not batch.failures
    indexes = [build_document_index(artifact) for artifact in batch.documents]
    validate_presentation_outline(
        outline,
        CorpusLookup.from_artifacts_indexes(batch.documents, indexes),
    )
    assert len(outline.all_slides()) == requirements.target_slide_count

    output_path = write_outline_json(outline, tmp_path / "outline.json")
    loaded = output_path.read_text(encoding="utf-8")
    assert "\"schema_version\"" in loaded
    assert type(outline).model_validate_json(loaded) == outline
