"""Offline end-to-end coverage for the semantic presentation-content boundary."""

from __future__ import annotations

import asyncio

from docx import Document

from presentation_pipeline import generate_plan, write_presentation_content_json
from presentation_pipeline.planning import PresentationRequirements
from presentation_pipeline.planning.models import EvidenceSelection, PresentationOutline
from presentation_pipeline.synthesis import (
    PresentationContent,
    SlideContent,
    build_presentation_content,
    build_slide_contexts,
    generate_slide_contents,
)
from presentation_pipeline.understanding.models import DocumentDigest
from presentation_pipeline.understanding.models import ChunkDigest
from presentation_pipeline.retrieval.models import LocalCandidateSelection


class _DeterministicPipelineGenerator:
    """Return schema-valid results derived only from each stage's supplied data."""

    async def generate(self, *, system_prompt, input_data, response_model):
        if response_model in (DocumentDigest, ChunkDigest):
            document = input_data["document"]
            if response_model is DocumentDigest and "fragments" in input_data:
                fragment = input_data["fragments"][0]
                return {
                    "doc_id": document["doc_id"],
                    "summary": fragment["summary"],
                    "topics": fragment["topics"],
                    "key_facts": fragment["key_facts"],
                }
            evidence = input_data["evidence"][0]
            result = {
                "doc_id": document["doc_id"],
                "summary": "Quarterly growth briefing",
                "key_facts": [
                    {
                        "claim": evidence.get("text", "Source-backed quarterly update"),
                        "evidence": [
                            {
                                "doc_id": document["doc_id"],
                                "evidence_ids": [evidence["evidence_id"]],
                            }
                        ],
                    }
                ],
            }
            if response_model is ChunkDigest:
                result["window_id"] = input_data["window"]["window_id"]
            return result
        if response_model is LocalCandidateSelection:
            evidence = input_data["evidence"][0]
            return {
                "candidates": [{
                    "doc_id": input_data["window"]["doc_id"],
                    "evidence_id": evidence["evidence_id"],
                    "reason": "Primary source for the briefing",
                }]
            }
        if response_model is EvidenceSelection:
            evidence = input_data["candidate_evidence"][0]
            return {
                "selected": [
                    {
                        "doc_id": evidence["doc_id"],
                        "evidence_id": evidence["evidence_id"],
                        "reason": "Primary source for the briefing",
                    }
                ]
            }
        if response_model is PresentationOutline:
            selected = input_data["selected_evidence"][0]
            return {
                "title": "Quarterly update",
                "objective": "Brief leadership",
                "narrative": "Introduce the update, then explain the evidence.",
                "sections": [
                    {
                        "section_id": "section-1",
                        "title": "Update",
                        "purpose": "Brief leadership",
                        "slides": [
                            {
                                "slide_id": "slide-1",
                                "title": "Quarterly update",
                                "purpose": "title",
                                "message": "Quarterly update",
                            },
                            {
                                "slide_id": "slide-2",
                                "title": "Evidence",
                                "purpose": "content",
                                "message": "The source records the quarterly update.",
                                "evidence": [
                                    {
                                        "doc_id": selected["doc_id"],
                                        "evidence_ids": [selected["evidence_id"]],
                                    }
                                ],
                            },
                        ],
                    }
                ],
            }
        if response_model is SlideContent:
            slide = input_data["slide"]
            if slide["purpose"] == "title":
                return {"slide_id": slide["slide_id"], "elements": []}
            evidence = input_data["evidence"][0]
            return {
                "slide_id": slide["slide_id"],
                "elements": [
                    {
                        "kind": "text",
                        "text": "The supplied document contains the quarterly update.",
                        "evidence": [
                            {
                                "doc_id": evidence["doc_id"],
                                "evidence_ids": [evidence["evidence_id"]],
                            }
                        ],
                    }
                ],
            }
        raise AssertionError(f"unexpected response model: {response_model}")


def test_real_docx_reaches_validated_semantic_content_json_offline(tmp_path) -> None:
    source_path = tmp_path / "briefing.docx"
    document = Document()
    document.add_heading("Quarterly update", level=1)
    document.add_paragraph("Revenue increased during the quarter.")
    document.save(source_path)

    requirements = PresentationRequirements(
        goal="Brief leadership",
        audience="Leadership team",
        target_slide_count=2,
        presentation_type="briefing",
    )
    generator = _DeterministicPipelineGenerator()

    async def run() -> tuple[PresentationContent, list[str]]:
        plan = await generate_plan(
            [source_path],
            requirements,
            generator,
            extraction_workers=1,
            llm_concurrency=1,
        )
        contexts = build_slide_contexts(plan)
        slides = await generate_slide_contents(contexts, generator, concurrency=1)
        return build_presentation_content(slides, contexts), [
            context.slide.slide_id for context in contexts
        ]

    content, expected_slide_ids = asyncio.run(run())
    assert [slide.slide_id for slide in content.slides] == expected_slide_ids == [
        "slide-1",
        "slide-2",
    ]

    output_path = write_presentation_content_json(content, tmp_path / "content.json")
    assert PresentationContent.model_validate_json(
        output_path.read_text(encoding="utf-8")
    ) == content
