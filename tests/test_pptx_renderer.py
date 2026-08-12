"""Offline contract and package tests for the presentation renderer."""

from __future__ import annotations

import base64

import pytest

from presentation_pipeline.common.references import EvidenceRef
from presentation_pipeline.planning.models import SlideOutline, SlidePurpose
from presentation_pipeline.rendering import (
    Box,
    LayoutArchetype,
    PresentationLayout,
    ResolvedElement,
    build_slide_layout,
    build_presentation_layout,
    render_presentation,
    resolve_render_inputs,
    select_layout,
    text_width_units,
    validate_layout,
)
from presentation_pipeline.synthesis.context import SlideContext
from presentation_pipeline.synthesis.models import (
    BulletItem,
    BulletListContent,
    ChartContent,
    SlideContent,
    TableContent,
    TextContent,
)


def _context(slide_id: str, *, purpose: SlidePurpose = SlidePurpose.CONTENT) -> SlideContext:
    return SlideContext(
        presentation_title="Policy briefing",
        presentation_objective="Support a decision",
        presentation_narrative="Evidence to action",
        section_id="section-1",
        section_title="Current position",
        section_purpose="Orient leaders",
        slide_index=1,
        total_slides=1,
        slide=SlideOutline(slide_id=slide_id, title="Decision in view", purpose=purpose, message="Move this quarter"),
    )


def _text_content(slide_id: str) -> SlideContent:
    return SlideContent(
        slide_id=slide_id,
        elements=[TextContent(kind="text", text="A practical policy decision is required now.", evidence=[EvidenceRef(doc_id="doc-1", evidence_ids=["ev-1"])])],
    )


def _write_pixel(path) -> None:
    path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScLvewAAAABJRU5ErkJggg=="))


def test_models_are_strict_and_geometry_is_validated() -> None:
    with pytest.raises(ValueError, match="extra"):
        Box(x=0.0, y=0.0, width=1.0, height=1.0, unknown=True)
    with pytest.raises(ValueError, match="greater than or equal"):
        Box(x=-0.1, y=0.0, width=1.0, height=1.0)
    with pytest.raises(ValueError, match="incomplete"):
        ResolvedElement(element_index=0, kind="chart")


def test_default_theme_is_traditional_chinese_first() -> None:
    from presentation_pipeline.rendering import ExecutivePolicyTheme

    theme = ExecutivePolicyTheme()
    assert theme.title_font == theme.body_font == theme.east_asian_font == "Microsoft JhengHei"
    assert theme.fallback_font == "Noto Sans TC"


def test_cjk_width_and_splitter_are_conservative_and_punctuation_aware() -> None:
    from presentation_pipeline.rendering.layout import _split_text

    assert text_width_units("漢字漢字") > text_width_units("abcd")
    chunks = _split_text("第一段說明，這裡有更多內容。第二段繼續說明，並且仍然需要分頁。", 12)
    assert len(chunks) > 1
    assert any(chunk.endswith("。") for chunk in chunks[:-1])
    assert not any(chunk.startswith(("，", "。", "！", "？", "；", "：", "、")) for chunk in chunks)


def test_layout_selection_prefers_decorative_image_and_places_columns() -> None:
    content = _text_content("visual")
    decorative = {"element_index": -1, "kind": "image", "path": "/tmp/decorative.png"}
    assert select_layout(_context("visual"), content, [decorative]) is LayoutArchetype.VISUAL_TEXT
    layout = build_slide_layout(_context("visual"), content, [decorative])
    assert layout.archetype is LayoutArchetype.VISUAL_TEXT
    assert [element.element_index for element in layout.elements] == [None, -1, 0]
    assert not [item for item in validate_layout(layout) if item.severity == "error"]


def test_text_and_list_select_two_columns_without_overlaps() -> None:
    content = SlideContent(
        slide_id="columns",
        elements=[
            TextContent(kind="text", text="Decision-makers need a clear baseline.", evidence=[EvidenceRef(doc_id="doc-1", evidence_ids=["ev-1"])]),
            BulletListContent(kind="list", items=[BulletItem(text="Set the owner", evidence=[EvidenceRef(doc_id="doc-1", evidence_ids=["ev-1"])])]),
        ],
    )
    layout = build_slide_layout(_context("columns"), content)
    assert layout.archetype is LayoutArchetype.TWO_COLUMN
    assert not [item for item in validate_layout(layout) if item.severity == "error"]


def test_renderer_creates_editable_table_chart_and_reopenable_package(tmp_path) -> None:
    pytest.importorskip("pptx")
    image_path = tmp_path / "visual.png"
    _write_pixel(image_path)
    visual = build_slide_layout(
        _context("visual"), _text_content("visual"), [{"element_index": -1, "kind": "image", "path": str(image_path)}]
    )
    table_content = SlideContent(slide_id="table", elements=[TableContent(kind="table", doc_id="doc-1", evidence_id="table-1")])
    table = build_slide_layout(
        _context("table"), table_content,
        [{"element_index": 0, "kind": "table", "rows": [["Owner", "Action"], ["Policy", "Approve"]], "header_rows": [0]}],
    )
    chart_content = SlideContent(slide_id="chart", elements=[ChartContent(kind="chart", doc_id="doc-1", evidence_id="chart-1", chart_type="bar")])
    chart = build_slide_layout(
        _context("chart"), chart_content,
        [{"element_index": 0, "kind": "chart", "chart_type": "bar", "categories": ["Now", "Target"], "series": [{"name": "Coverage", "values": [42.0, 75.0]}]}],
    )
    output = tmp_path / "briefing.pptx"
    report = render_presentation(PresentationLayout(slides=[visual, table, chart]), output)
    assert report.verified and report.slides_rendered == 3 and output.is_file()

    from pptx import Presentation

    deck = Presentation(output)
    assert len(deck.slides) == 3
    assert any(shape.has_table for shape in deck.slides[1].shapes)
    assert any(shape.has_chart for shape in deck.slides[2].shapes)


def test_resolver_keeps_table_chart_fallbacks_and_rejects_unsafe_images(tmp_path) -> None:
    table = resolve_render_inputs(
        "table",
        [TableContent(kind="table", doc_id="doc-1", evidence_id="table-1", title="Owners")],
        source_lookup={("doc-1", "table-1"): {"text": "Owner and action source table."}},
    )
    assert table.elements[0].kind == "text"
    assert table.elements[0].source_attribution is not None
    assert table.diagnostics[0].code == "TABLE_FALLBACK"

    chart = resolve_render_inputs(
        "chart",
        [ChartContent(kind="chart", doc_id="doc-1", evidence_id="chart-1", chart_type="bar")],
        source_lookup={("doc-1", "chart-1"): {"text": "Coverage chart source."}},
    )
    assert chart.elements[0].kind == "text"
    assert chart.diagnostics[0].code == "NATIVE_CHART_VALUES_UNAVAILABLE"

    unsafe = resolve_render_inputs(
        "image",
        [{"kind": "image", "doc_id": "doc-1", "evidence_id": "image-1", "path": "/tmp/not-approved.png"}],
        asset_root=tmp_path,
    )
    assert unsafe.elements[0].kind == "text"
    assert unsafe.diagnostics[0].code == "SOURCE_IMAGE_PATH_UNSAFE"

    exact_chart = resolve_render_inputs(
        "chart-data",
        [ChartContent(kind="chart", doc_id="doc-1", evidence_id="chart-data-1", chart_type="bar", title="Coverage")],
        source_lookup={("doc-1", "chart-data-1"): {"structured_data": {
            "categories": ["Now", "Target"],
            "series": [{"name": "Coverage", "column": 1, "values": [42, 75]}],
        }}},
        source_filenames={"doc-1": "policy.docx"},
    )
    assert exact_chart.elements[0].kind == "chart"
    assert exact_chart.elements[0].series[0].values == [42.0, 75.0]
    assert exact_chart.elements[0].source_attribution.filename == "policy.docx"

    multi_source = resolve_render_inputs(
        "multi-source",
        [TextContent(kind="text", text="Cross-document conclusion", evidence=[
            EvidenceRef(doc_id="doc-1", evidence_ids=["ev-1", "ev-2"]),
            EvidenceRef(doc_id="doc-2", evidence_ids=["ev-3"]),
        ])],
        source_filenames={"doc-1": "policy.docx", "doc-2": "access.docx"},
    )
    assert [(item.filename, item.evidence_ids) for item in multi_source.elements[0].attribution] == [
        ("policy.docx", ["ev-1", "ev-2"]),
        ("access.docx", ["ev-3"]),
    ]


def test_layout_paginates_wide_tables_and_multiple_images_with_stable_ids(tmp_path) -> None:
    rows = [[f"H{column}" for column in range(7)]] + [[f"{row}-{column}" for column in range(7)] for row in range(10)]
    table_content = SlideContent(slide_id="wide", elements=[TableContent(kind="table", doc_id="doc-1", evidence_id="table-1")])
    layouts = build_presentation_layout(
        [_context("wide")],
        [table_content],
        {"wide": [{"element_index": 0, "kind": "table", "rows": rows, "header_rows": [0]}]},
    )
    assert [slide.slide_id for slide in layouts.slides] == ["wide--p01", "wide--p02", "wide--p03", "wide--p04"]
    assert all(slide.semantic_slide_id == "wide" for slide in layouts.slides)
    assert layouts.slides[1].physical_slide_id == "wide--p02"
    assert layouts.slides[1].continuation_index == 1 and layouts.slides[1].is_continuation
    assert any(element.text == "Decision in view (continued)" for element in layouts.slides[1].elements)
    first_table = next(element.payload.table_payload for element in layouts.slides[0].elements if element.kind == "table")
    repeated_key_table = next(element.payload.table_payload for element in layouts.slides[2].elements if element.kind == "table")
    assert len(first_table.rows) == 9  # repeated header plus eight body rows
    assert repeated_key_table.plain_rows()[0] == ["H0", "H6"]

    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    _write_pixel(first)
    _write_pixel(second)
    content = _text_content("images")
    image_layouts = build_presentation_layout(
        [_context("images")], [content],
        {"images": [
            {"element_index": -1, "kind": "image", "path": str(first)},
            {"element_index": 1, "kind": "image", "path": str(second)},
        ]},
    )
    assert [slide.slide_id for slide in image_layouts.slides] == ["images--p01", "images--p02"]


def test_renderer_writes_source_footer_real_bullets_and_east_asian_font(tmp_path) -> None:
    pytest.importorskip("pptx")
    content = SlideContent(
        slide_id="sources",
        elements=[BulletListContent(kind="list", items=[BulletItem(text="繁體中文 bullet", evidence=[EvidenceRef(doc_id="doc-1", evidence_ids=["ev-1"])])])],
    )
    layout = build_slide_layout(
        _context("sources"), content,
        [{"element_index": 0, "kind": "list", "items": ["繁體中文 bullet"], "source_attribution": {"doc_id": "doc-1", "evidence_ids": ["ev-1"], "filename": "policy.docx"}}],
    )
    output = tmp_path / "sources.pptx"
    report = render_presentation(PresentationLayout(slides=[layout]), output)
    assert report.source_footer_count == 1
    from pptx import Presentation

    slide = Presentation(output).slides[0]
    xml = "".join(shape.element.xml for shape in slide.shapes)
    assert "a:buChar" in xml
    assert "Microsoft JhengHei" in xml
    slide_text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
    assert "來源：policy.docx" in slide_text
    assert "ev-1" not in slide_text
    assert 'b="1"' in xml


def test_source_footer_is_compact_and_localized_for_chinese_multi_source() -> None:
    content = _text_content("sources-zh")
    layout = build_slide_layout(
        _context("sources-zh"), content,
        [{"element_index": 0, "kind": "text", "text": "繁體中文重點", "attribution": [
            {"doc_id": "doc-1", "evidence_ids": ["ev-1"]},
            {"doc_id": "doc-2", "evidence_ids": ["ev-2"]},
            {"doc_id": "doc-3", "evidence_ids": ["ev-3"]},
        ]}],
    )
    footer = next(item.text for item in layout.elements if item.role == "footer")
    # Without filenames these logical source IDs still use Chinese UI because
    # the slide content is Traditional Chinese.
    assert footer == "來源：doc-1；doc-2（另 1 份）"
