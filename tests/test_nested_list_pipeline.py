"""End-to-end regression coverage for nested DOCX list evidence."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_pipeline.api import extract_docx
from presentation_pipeline.corpus.models import DocumentArtifact
from presentation_pipeline.indexing import EvidenceKind, build_document_index
from presentation_pipeline.planning.models import EvidenceRef
from presentation_pipeline.validation import CorpusLookup


def _set_numbering(paragraph: object, *, num_id: int, ilvl: int) -> None:
    """Write explicit numbering XML so this is a real nested DOCX list."""
    properties = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(ilvl))
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    numbering.extend((level, number))
    properties.append(numbering)


def _nested_list_docx(tmp_path: Path) -> Path:
    document = Document()
    parent = document.add_paragraph("Parent item", style="List Number")
    child = document.add_paragraph("Nested child", style="List Number 2")
    sibling = document.add_paragraph("Second parent", style="List Number")
    for paragraph, level in ((parent, 0), (child, 1), (sibling, 0)):
        _set_numbering(paragraph, num_id=5, ilvl=level)
    path = tmp_path / "nested-list.docx"
    document.save(path)
    return path


def test_nested_docx_list_survives_extraction_indexing_and_provenance(tmp_path: Path) -> None:
    """Nested child arrays must remain JSON-safe through the planning boundary."""
    extraction = extract_docx(_nested_list_docx(tmp_path))
    artifact = DocumentArtifact(
        job_id="nested-list-regression",
        doc_id=extraction.document.doc_id,
        filename=extraction.document.filename,
        extraction=extraction,
    )

    index = build_document_index(artifact)
    list_evidence = next(item for item in index.evidence if item.kind is EvidenceKind.LIST)
    root_item = list_evidence.structured_data["items"][0]
    child_item = root_item["children"][0]

    assert list_evidence.text == "Parent item\nNested child\nSecond parent"
    assert child_item["children"] == []
    assert child_item["level"] == 1
    assert root_item["source_node_id"] in list_evidence.source_node_ids
    assert child_item["source_node_id"] in list_evidence.source_node_ids
    assert child_item["source_node_id"] in {
        node.node_id for node in extraction.nodes if node.kind.value == "paragraph"
    }

    # The serialized index is suitable for provider prompts/artifacts, and the
    # lookup validates every evidence-to-source relationship without repair.
    json.dumps(index.model_dump(mode="json"), ensure_ascii=False)
    lookup = CorpusLookup.from_artifacts_indexes([artifact], [index])
    lookup.validate_ref(EvidenceRef(doc_id=artifact.doc_id, evidence_ids=[list_evidence.evidence_id]))
