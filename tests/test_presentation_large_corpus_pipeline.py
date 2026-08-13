"""Offline scale regression: all planning requests remain bounded."""

from __future__ import annotations

import asyncio

from docx import Document

from presentation_pipeline import PlanningScaleConfig, generate_plan
from presentation_pipeline.budgeting import (
    InputBudget,
    PlanningBudgets,
    RetrievalBudget,
    UnderstandingBudgets,
    Utf8ByteTokenEstimator,
    estimate_request_tokens,
)
from presentation_pipeline.planning import PresentationRequirements
from presentation_pipeline.planning.models import EvidenceSelection, PresentationOutline
from presentation_pipeline.retrieval.models import (
    CandidateReductionSelection,
    LocalCandidateSelection,
)
from presentation_pipeline.understanding.models import ChunkDigest, DigestFragment, DocumentDigest


class _BoundedLargeCorpusGenerator:
    """A deterministic provider double that records every bounded request."""

    def __init__(self) -> None:
        self.calls: list[tuple[str, dict[str, object], type[object]]] = []
        self.active = 0
        self.maximum_active = 0

    async def generate(self, *, system_prompt, input_data, response_model):
        self.calls.append((system_prompt, input_data, response_model))
        self.active += 1
        self.maximum_active = max(self.maximum_active, self.active)
        await asyncio.sleep(0)
        self.active -= 1
        if response_model is ChunkDigest:
            item = input_data["evidence"][0]
            return {
                "doc_id": input_data["document"]["doc_id"],
                "window_id": input_data["window"]["window_id"],
                "summary": "bounded chunk",
                "key_facts": [{"claim": "source fact", "evidence": [{"doc_id": input_data["document"]["doc_id"], "evidence_ids": [item["evidence_id"]]}]}],
            }
        if response_model is DocumentDigest:
            fragment = input_data["fragments"][0]
            return {"doc_id": input_data["document"]["doc_id"], "summary": "bounded digest", "topics": [], "key_facts": fragment["key_facts"]}
        if response_model is DigestFragment:
            fragment = input_data["fragments"][0]
            return {"doc_id": input_data["document"]["doc_id"], "fragment_id": "reduced", "summary": "reduced", "topics": [], "key_facts": fragment["key_facts"]}
        if response_model is CandidateReductionSelection:
            return {
                "candidates": [
                    {
                        "doc_id": item["doc_id"],
                        "evidence_id": item["evidence_id"],
                    }
                    for item in input_data["candidate_descriptors"]
                ]
            }
        if response_model is LocalCandidateSelection:
            return {"candidates": [
                {"doc_id": input_data["window"]["doc_id"], "evidence_id": item["evidence_id"], "reason": "local"}
                for item in input_data["evidence"][:2]
            ]}
        if response_model is EvidenceSelection:
            item = input_data["candidate_evidence"][0]
            return {"selected": [{"doc_id": item["doc_id"], "evidence_id": item["evidence_id"], "reason": "selected"}]}
        if response_model is PresentationOutline:
            return {"title": "Bounded", "objective": "Brief", "narrative": "Source-backed", "sections": [{"section_id": "s", "title": "S", "purpose": "brief", "slides": [{"slide_id": "slide-1", "title": "Title", "purpose": "title", "message": "Bounded"}]}]}
        raise AssertionError(response_model)


def _scale() -> PlanningScaleConfig:
    return PlanningScaleConfig(
        token_counter=Utf8ByteTokenEstimator(),
        budgets=PlanningBudgets(
            understanding=UnderstandingBudgets(window=InputBudget(1_500), reduction=InputBudget(1_200)),
            retrieval=RetrievalBudget(
                request_budget=InputBudget(1_800), reduction_budget=InputBudget(1_800),
                selection_budget=InputBudget(1_800), max_candidates_per_window=2,
                max_global_candidates=2,
            ),
            outline_generation=InputBudget(1_400),
        ),
    )


def test_large_synthetic_docx_uses_only_bounded_requests_and_reduces_candidates(tmp_path) -> None:
    source = tmp_path / "large.docx"
    document = Document()
    document.add_heading("Large corpus", level=1)
    table = document.add_table(rows=1, cols=3)
    for column, heading in zip(table.rows[0].cells, ("Policy", "Value", "Notes")):
        column.text = heading
    for number in range(80):
        cells = table.add_row().cells
        cells[0].text = f"Policy {number}"
        cells[1].text = f"Value {number}"
        cells[2].text = "oversized table evidence " * 12
    for number in range(24):
        document.add_paragraph(f"Finding {number}: " + ("evidence text " * 16))
    document.save(source)
    scale = _scale()
    generator = _BoundedLargeCorpusGenerator()
    result = asyncio.run(generate_plan(
        [source], PresentationRequirements(goal="Brief", audience="Leaders", target_slide_count=1),
        generator, extraction_workers=1, llm_concurrency=2, scale_config=scale,
    ))

    assert result.candidates is not None
    assert len(result.candidates.candidates) <= 2
    assert result.outline.all_slides()[0].slide_id == "slide-1"
    assert generator.maximum_active <= 2
    reduction_calls = [
        payload
        for _, payload, model in generator.calls
        if model is CandidateReductionSelection
    ]
    discovery_calls = [payload for _, payload, model in generator.calls if model is LocalCandidateSelection and "evidence" in payload]
    assert len(discovery_calls) > 1
    assert len(reduction_calls) >= 2  # at least two levels/groups of reduction
    assert all("candidate_descriptors" in payload for payload in reduction_calls)
    assert all(
        set(descriptor)
        == {"doc_id", "evidence_id", "kind", "reason", "score", "hit_count", "slice_count"}
        for payload in reduction_calls
        for descriptor in payload["candidate_descriptors"]
    )
    assert all("evidence_catalogue" not in payload for _, payload, _ in generator.calls)

    counter = scale.token_counter
    for prompt, payload, model in generator.calls:
        if model is ChunkDigest:
            budget = scale.budgets.understanding.window
        elif model in (DocumentDigest, DigestFragment):
            budget = scale.budgets.understanding.reduction
        elif model is CandidateReductionSelection:
            budget = scale.budgets.retrieval.reduction_budget
        elif model is LocalCandidateSelection:
            budget = scale.budgets.retrieval.request_budget
        elif model is EvidenceSelection:
            budget = scale.budgets.retrieval.selection_budget
        else:
            budget = scale.budgets.outline_generation
        assert estimate_request_tokens(prompt, payload, counter) <= budget.usable_input_tokens
