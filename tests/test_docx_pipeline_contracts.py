"""Cross-cutting contract tests for package safety, inline features, and CLI."""

from __future__ import annotations

import base64
from io import BytesIO
import json
from pathlib import Path
import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from pydantic import ValidationError

from docx_pipeline.api import extract_docx
from docx_pipeline.cli import _run
from docx_pipeline.ir import ExtractionResult, SourceNode
from docx_pipeline.package import DocxPackage
from docx_pipeline.relationships import parse_relationships, resolve_relationship_target


PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4"
    "z8DwHwAFgAI/ScL9VwAAAABJRU5ErkJggg=="
)
VML_NS = "urn:schemas-microsoft-com:vml"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _add_explicit_hyperlink(paragraph: object, text: str, url: str) -> None:
    part = paragraph.part  # type: ignore[attr-defined]
    relationship_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)  # type: ignore[attr-defined]


def _add_field_hyperlink(paragraph: object, text: str, url: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f' HYPERLINK "{url}" ')
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    field.append(run)
    paragraph._p.append(field)  # type: ignore[attr-defined]


def _add_inline_sdt(paragraph: object, text: str) -> None:
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "transparent fixture")
    properties.append(alias)
    contents = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    contents.append(run)
    control.extend((properties, contents))
    paragraph._p.append(control)  # type: ignore[attr-defined]


def _add_text_box(paragraph: object, text: str) -> None:
    """Build the minimal DrawingML-shaped textbox consumed by the XML walker."""
    drawing = OxmlElement("w:drawing")
    content = OxmlElement("w:txbxContent")
    boxed_paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    boxed_paragraph.append(run)
    content.append(boxed_paragraph)
    drawing.append(content)
    paragraph._p.append(drawing)  # type: ignore[attr-defined]


def _textbox_content(text: str) -> object:
    content = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    paragraph.append(run)
    content.append(paragraph)
    return content


def _add_vml_text_box(paragraph: object, text: str) -> None:
    pict = OxmlElement("w:pict")
    shape = etree.Element(f"{{{VML_NS}}}shape", nsmap={"v": VML_NS})
    textbox = etree.Element(f"{{{VML_NS}}}textbox")
    textbox.append(_textbox_content(text))
    shape.append(textbox)
    pict.append(shape)
    paragraph._p.append(pict)  # type: ignore[attr-defined]


def _add_mc_text_box(paragraph: object, text: str, *, alias: str = "wps") -> None:
    alternate = etree.Element(f"{{{MC_NS}}}AlternateContent", nsmap={"mc": MC_NS})
    shape_ns = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    choice = etree.Element(
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice",
        nsmap={alias: shape_ns},
    )
    choice.set("Requires", alias)
    shape = etree.Element(f"{{{shape_ns}}}wsp")
    textbox = etree.Element(f"{{{shape_ns}}}txbx")
    textbox.append(_textbox_content(text))
    shape.append(textbox)
    choice.append(shape)
    fallback = etree.Element(f"{{{MC_NS}}}Fallback")
    fallback_shape = etree.Element(f"{{{VML_NS}}}shape", nsmap={"v": VML_NS})
    fallback_textbox = etree.Element(f"{{{VML_NS}}}textbox")
    fallback_textbox.append(_textbox_content(text))
    fallback_shape.append(fallback_textbox)
    fallback_pict = OxmlElement("w:pict")
    fallback_pict.append(fallback_shape)
    fallback.append(fallback_pict)
    alternate.extend((choice, fallback))
    paragraph._p.append(alternate)  # type: ignore[attr-defined]


def _paragraph_element(text: str) -> object:
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    token = OxmlElement("w:t")
    token.text = text
    run.append(token)
    paragraph.append(run)
    return paragraph


def _add_mc_block(
    document: Document,
    choice_text: str,
    fallback_text: str | None,
    *,
    supported_choice: bool = True,
) -> None:
    shape_ns = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    alternate = etree.Element(f"{{{MC_NS}}}AlternateContent", nsmap={"mc": MC_NS})
    prefix = "wps" if supported_choice else "unsupported"
    namespace = shape_ns if supported_choice else "urn:test:unsupported"
    choice = etree.Element(f"{{{MC_NS}}}Choice", nsmap={prefix: namespace})
    choice.set("Requires", prefix)
    choice.set(f"{{{namespace}}}marker", "1")
    choice.append(_paragraph_element(choice_text))
    alternate.append(choice)
    if fallback_text is not None:
        fallback = etree.Element(f"{{{MC_NS}}}Fallback")
        fallback.append(_paragraph_element(fallback_text))
        alternate.append(fallback)
    body = document._body._element
    body.insert(len(body) - 1, alternate)


def _add_nested_mc_image(document: Document, image_path: Path) -> None:
    relationship_id, _ = document.part.get_or_add_image(str(image_path))
    paragraph = document.add_paragraph()
    drawing = OxmlElement("w:drawing")
    alternate = etree.Element(f"{{{MC_NS}}}AlternateContent", nsmap={"mc": MC_NS})
    shape_ns = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    choice = etree.Element(f"{{{MC_NS}}}Choice", nsmap={"shape": shape_ns})
    choice.set("Requires", "shape")
    choice_blip = etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    choice_blip.set(qn("r:embed"), relationship_id)
    choice.append(choice_blip)
    fallback = etree.Element(f"{{{MC_NS}}}Fallback")
    fallback_blip = etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    fallback_blip.set(qn("r:embed"), relationship_id)
    fallback.append(fallback_blip)
    alternate.extend((choice, fallback))
    drawing.append(alternate)
    paragraph._p.append(drawing)


def _save(document: Document, path: Path) -> Path:
    document.save(path)
    return path


def _nodes(result: object, kind: str) -> list[object]:
    return [node for node in result.nodes if node.kind.value == kind]  # type: ignore[attr-defined]


def _text(nodes: list[object]) -> str:
    return "".join(str(node.payload.get("text", "")) for node in nodes)  # type: ignore[attr-defined]


def _replace_zip_member(path: Path, member: str, content: bytes) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            destination.writestr(info, content if info.filename == member else source.read(info.filename))
    return output.getvalue()


def test_relationship_targets_are_owner_relative_and_cannot_escape_package() -> None:
    assert resolve_relationship_target("word/header1.xml", "media/image1.png") == "word/media/image1.png"
    assert resolve_relationship_target("word/header1.xml", "../media/image1.png") == "media/image1.png"
    with pytest.raises(ValueError, match="escapes package root"):
        resolve_relationship_target("word/document.xml", "../../outside.xml")
    with pytest.raises(ValueError, match="absolute"):
        resolve_relationship_target("word/document.xml", "/word/media/image1.png")

    relationships = parse_relationships(
        b'''<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
        <Relationship Id="rId1" Type="image" Target="media/a.png"/>
        </Relationships>''',
        "word/header1.xml",
    )
    assert relationships["rId1"].target_part_name == "word/media/a.png"


def test_preflight_diagnoses_unsafe_relationship_without_mutating_source(tmp_path: Path) -> None:
    path = _save(Document(), tmp_path / "unsafe.docx")
    relationships = b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rIdUnsafe" Type="urn:test" Target="../../outside.xml"/>
    </Relationships>'''
    package = DocxPackage.from_bytes(_replace_zip_member(path, "word/_rels/document.xml.rels", relationships))
    before = package.original_bytes
    preflight = package.preflight(sanitize=True)
    assert package.original_bytes == before
    assert preflight.effective_bytes == before
    assert preflight.original_sha256 == preflight.effective_sha256
    assert any(item.code == "UNSAFE_RELATIONSHIP_TARGET" and item.relationship_id == "rIdUnsafe" for item in preflight.diagnostics)
    assert any(item.code == "SANITIZATION_NOT_APPLIED" for item in preflight.diagnostics)


def test_inline_sdt_is_a_container_without_hiding_visible_text(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("before ")
    _add_inline_sdt(paragraph, "inside")
    paragraph.add_run(" after")
    result = extract_docx(_save(document, tmp_path / "inline-sdt.docx"))

    controls = _nodes(result, "content_control")
    assert len(controls) == 1
    assert controls[0].payload["inline"] is True
    text_nodes = _nodes(result, "text_run")
    assert _text(text_nodes) == "before inside after"
    assert any(node.parent_id == controls[0].node_id and node.payload["text"] == "inside" for node in text_nodes)
    assert result.inventory.totals["structured_document_tags"] == 1
    assert result.coverage.metrics["structured_document_tags"].extracted == 1


def test_explicit_and_field_hyperlinks_are_both_accounted_for(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("A ")
    _add_explicit_hyperlink(paragraph, "explicit", "https://example.test/explicit")
    paragraph.add_run(" B ")
    _add_field_hyperlink(paragraph, "field", "https://example.test/field")
    result = extract_docx(_save(document, tmp_path / "links.docx"))

    hyperlinks = _nodes(result, "hyperlink")
    fields = _nodes(result, "field")
    assert len(hyperlinks) == 1
    assert hyperlinks[0].payload["url"] == "https://example.test/explicit"
    assert len(fields) == 1
    assert "HYPERLINK" in fields[0].payload["instruction"]
    assert _text(_nodes(result, "text_run")) == "A explicit B field"
    assert result.inventory.totals["hyperlinks"] == 1
    assert result.inventory.totals["field_hyperlinks"] == 1
    assert result.coverage.metrics["hyperlinks"].source == 1
    assert result.coverage.metrics["hyperlinks"].extracted == 1


def test_text_box_is_traversed_as_a_typed_container(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("outside ")
    _add_text_box(paragraph, "boxed")
    paragraph.add_run(" tail")
    result = extract_docx(_save(document, tmp_path / "textbox.docx"))

    boxes = _nodes(result, "text_box")
    assert len(boxes) == 1
    text_nodes = _nodes(result, "text_run")
    assert _text(text_nodes) == "outside boxed tail"
    assert any(node.parent_id == boxes[0].node_id for node in result.nodes)
    assert result.inventory.totals["textboxes"] == 1
    assert result.coverage.metrics["textboxes"].extracted == 1


def test_vml_text_box_is_typed_and_conserves_visible_text(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("before ")
    _add_vml_text_box(paragraph, "inside VML")
    paragraph.add_run(" after")

    result = extract_docx(_save(document, tmp_path / "vml-textbox.docx"))

    boxes = _nodes(result, "text_box")
    assert len(boxes) == 1
    assert boxes[0].payload["placement"] == "vml"
    assert _text(_nodes(result, "text_run")) == "before inside VML after"
    assert result.inventory.totals["textboxes"] == 1
    assert result.coverage.silent_losses == 0
    assert not any(item.code == "UNSUPPORTED_VML_CONTENT" for item in result.diagnostics)


def test_mc_choice_textbox_excludes_duplicate_vml_fallback(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("before ")
    _add_mc_text_box(paragraph, "one textbox")
    paragraph.add_run(" after")

    result = extract_docx(_save(document, tmp_path / "mc-textbox.docx"))

    assert len(_nodes(result, "text_box")) == 1
    assert _text(_nodes(result, "text_run")) == "before one textbox after"
    assert result.inventory.totals["textboxes"] == 1
    assert result.coverage.silent_losses == 0


def test_mc_choice_support_is_based_on_namespace_uri_not_prefix(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    _add_mc_text_box(paragraph, "aliased textbox", alias="shape")

    result = extract_docx(_save(document, tmp_path / "mc-alias-textbox.docx"))

    assert len(_nodes(result, "text_box")) == 1
    assert _text(_nodes(result, "text_run")) == "aliased textbox"
    assert result.coverage.silent_losses == 0


def test_mc_block_choice_excludes_fallback_and_conserves_text(tmp_path: Path) -> None:
    document = Document()
    _add_mc_block(document, "selected block", "fallback block")

    result = extract_docx(_save(document, tmp_path / "mc-block.docx"))

    assert _text(_nodes(result, "text_run")) == "selected block"
    assert result.inventory.totals["visible_text_nodes"] == 1
    assert result.coverage.silent_losses == 0


def test_mc_block_without_effective_branch_is_diagnosed(tmp_path: Path) -> None:
    document = Document()
    _add_mc_block(document, "unavailable block", None, supported_choice=False)

    result = extract_docx(_save(document, tmp_path / "mc-unsupported-block.docx"))

    assert not _nodes(result, "text_run")
    assert any(item.code == "UNSUPPORTED_MARKUP_COMPATIBILITY_BRANCH" for item in result.diagnostics)
    assert len(_nodes(result, "unsupported")) == 1


def test_nested_mc_drawing_uses_only_selected_branch(tmp_path: Path) -> None:
    image = tmp_path / "pixel.png"
    image.write_bytes(PIXEL_PNG)
    document = Document()
    _add_nested_mc_image(document, image)

    result = extract_docx(_save(document, tmp_path / "nested-mc-drawing.docx"))

    assert len(_nodes(result, "image")) == 1
    assert result.inventory.totals["image_blips"] == 1
    assert result.coverage.silent_losses == 0


def test_vml_image_remains_an_image_occurrence(tmp_path: Path) -> None:
    image = tmp_path / "pixel.png"
    image.write_bytes(PIXEL_PNG)
    document = Document()
    relationship_id, _ = document.part.get_or_add_image(str(image))
    paragraph = document.add_paragraph()
    pict = OxmlElement("w:pict")
    shape = etree.Element(f"{{{VML_NS}}}shape", nsmap={"v": VML_NS})
    image_data = etree.Element(f"{{{VML_NS}}}imagedata")
    image_data.set(qn("r:id"), relationship_id)
    shape.append(image_data)
    pict.append(shape)
    paragraph._p.append(pict)

    result = extract_docx(_save(document, tmp_path / "vml-image.docx"))

    images = _nodes(result, "image")
    assert len(images) == 1
    assert images[0].payload["placement"] == "vml"
    assert len(result.assets) == 1
    assert result.coverage.metrics["image_blips"].source == 1
    assert result.coverage.metrics["image_blips"].extracted == 1
    assert result.coverage.silent_losses == 0
    assert not any(item.code in {"IMAGE_WITHOUT_RELATIONSHIP", "UNRESOLVED_IMAGE_RELATIONSHIP"} for item in result.diagnostics)


def test_coverage_reports_conservation_for_mixed_inline_fixture(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph("visible ")
    _add_inline_sdt(paragraph, "controlled")
    _add_explicit_hyperlink(paragraph, " link", "https://example.test")
    result = extract_docx(_save(document, tmp_path / "coverage.docx"))

    assert result.coverage.silent_losses == 0
    for feature in ("visible_text_nodes", "visible_text_chars", "hyperlinks", "structured_document_tags"):
        metric = result.coverage.metrics[feature]
        assert metric.source == metric.extracted
        assert metric.diagnosed == 0
        assert metric.ratio == 1.0


def test_heading_styles_are_resolved_into_replaceable_section_views(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("body")
    document.add_heading("Details", level=3)
    result = extract_docx(_save(document, tmp_path / "headings.docx"))

    heading_nodes = [
        node
        for node in result.nodes
        if node.kind.value == "paragraph" and node.payload.get("outline_level") is not None
    ]
    assert [(node.payload["style_name"], node.payload["outline_level"]) for node in heading_nodes] == [
        ("heading 1", 0),
        ("heading 3", 2),
    ]
    assert [(section.title, section.level) for section in result.views.sections] == [
        ("Introduction", 1),
        ("Details", 3),
    ]
    assert all(node.parent_id is not None for node in heading_nodes)


def test_cli_extract_is_canonical_and_validate_accepts_it(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("CLI fixture")
    source = _save(document, tmp_path / "cli.docx")
    first = tmp_path / "first.json"
    second = tmp_path / "second.json"
    assert _run(["extract", str(source), "--output", str(first)]) == 0
    assert _run(["extract", str(source), "--output", str(second)]) == 0
    assert first.read_bytes() == second.read_bytes()
    parsed = json.loads(first.read_text(encoding="utf-8"))
    assert parsed["schema_version"] == "1.0.0"
    assert _run(["validate", str(first)]) == 0


def test_ir_rejects_unknown_parent_invalid_status_and_schema_fields(tmp_path: Path) -> None:
    result = extract_docx(_save(Document(), tmp_path / "empty.docx"))
    payload = result.model_dump(mode="json")

    invalid_parent = json.loads(json.dumps(payload))
    invalid_parent["nodes"][0]["parent_id"] = "node-does-not-exist"
    with pytest.raises(ValidationError, match="unknown parent_id"):
        ExtractionResult.model_validate(invalid_parent)

    invalid_status = json.loads(json.dumps(payload))
    invalid_status["nodes"][0]["status"] = "silently_dropped"
    with pytest.raises(ValidationError, match="status"):
        ExtractionResult.model_validate(invalid_status)

    node = result.nodes[0].model_dump(mode="json")
    node["unrecognised_contract_field"] = True
    with pytest.raises(ValidationError, match="unrecognised_contract_field"):
        SourceNode.model_validate(node)
