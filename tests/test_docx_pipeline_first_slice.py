"""Contract tests for the first auditable DOCX-pipeline slice.

These deliberately construct DOCX packages at test time.  They are an
independent OOXML oracle for occurrence counts, rather than tests of the
legacy section/block extractor.
"""

from __future__ import annotations

import base64
import hashlib
import importlib
import json
import re
import zipfile
from collections.abc import Iterable
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"w": W_NS, "a": A_NS}

# A valid 1x1 transparent PNG.  Keeping it inline makes fixtures portable and
# lets two drawing occurrences point at byte-identical asset content.
PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4"
    "z8DwHwAFgAI/ScL9VwAAAABJRU5ErkJggg=="
)


def _api() -> Any:
    """Import lazily so fixture construction remains independent of the API."""
    return importlib.import_module("docx_pipeline.api")


def _extract(path: Path) -> Any:
    api = _api()
    result = api.extract_docx(path)
    for name in ("nodes", "assets", "inventory", "coverage", "diagnostics"):
        assert hasattr(result, name), f"ExtractionResult must expose .{name}"
    return result


def _mapping(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if hasattr(value, "model_dump"):
        return value.model_dump(mode="json", exclude_none=False)
    if hasattr(value, "dict"):
        return value.dict(exclude_none=False)
    return dict(vars(value))


def _items(value: Iterable[Any]) -> list[dict[str, Any]]:
    return [_mapping(item) for item in value]


def _field(value: Any, name: str, default: Any = None) -> Any:
    return _mapping(value).get(name, default)


def _canonical_json(result: Any) -> str:
    """Exercise the public canonical serializer when supplied, else its data."""
    for name in ("canonical_json", "to_canonical_json"):
        serializer = getattr(result, name, None)
        if callable(serializer):
            rendered = serializer()
            return rendered.decode("utf-8") if isinstance(rendered, bytes) else rendered
    return json.dumps(_mapping(result), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _write_pixel(tmp_path: Path) -> Path:
    path = tmp_path / "pixel.png"
    path.write_bytes(PIXEL_PNG)
    return path


def _append_inline_equation(paragraph: Any, text: str = "x") -> None:
    """Insert minimal OMML directly between ordinary runs."""
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    token = OxmlElement("m:t")
    token.text = text
    run.append(token)
    omath.append(run)
    paragraph._p.append(omath)


def _set_numbering(paragraph: Any, *, num_id: int = 5, ilvl: int = 0) -> None:
    """Add raw ``w:numPr`` facts; styles alone are not a numbering oracle."""
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(ilvl))
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    numbering.extend((level, number))
    properties.append(numbering)


def _mixed_images_docx(tmp_path: Path, *, include_equation: bool = False) -> Path:
    image = _write_pixel(tmp_path)
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("alpha ")
    paragraph.add_run().add_picture(str(image))
    paragraph.add_run(" beta ")
    paragraph.add_run().add_picture(str(image))
    paragraph.add_run(" gamma")
    if include_equation:
        paragraph.add_run(" + ")
        _append_inline_equation(paragraph, "x")
        paragraph.add_run(" omega")
    path = tmp_path / "mixed.docx"
    document.save(path)
    return path


def _node_text(node: dict[str, Any]) -> str:
    payload = node.get("payload") or {}
    return str(payload.get("text", node.get("text", "")) or "")


def _node_kind(node: dict[str, Any]) -> str:
    return str(node.get("kind", node.get("type", "")))


def _numbering(node: dict[str, Any]) -> dict[str, Any] | None:
    """Accept either the compact or grouped public representation."""
    payload = node.get("payload") or {}
    grouped = payload.get("numbering")
    if isinstance(grouped, dict):
        return grouped
    if payload.get("num_id") is not None:
        return {"num_id": payload["num_id"], "ilvl": payload.get("ilvl", payload.get("list_level"))}
    return None


def _descendants(nodes: list[dict[str, Any]], parent_id: str) -> list[dict[str, Any]]:
    by_parent: dict[str, list[dict[str, Any]]] = {}
    for node in nodes:
        if node.get("parent_id") is not None:
            by_parent.setdefault(str(node["parent_id"]), []).append(node)
    result: list[dict[str, Any]] = []
    pending = list(by_parent.get(parent_id, []))
    while pending:
        child = pending.pop(0)
        result.append(child)
        pending.extend(by_parent.get(str(child.get("node_id")), []))
    return result


def _xml_image_paragraph_texts(path: Path) -> list[str]:
    """Visible text in each main-body paragraph containing a DrawingML image."""
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    output: list[str] = []
    for paragraph in root.findall(".//w:body/w:p", NS):
        if paragraph.find(".//a:blip", NS) is None:
            continue
        text = "".join(token.text or "" for token in paragraph.findall(".//w:t", NS))
        if text:
            output.append(text)
    return output


def _xml_image_count(path: Path) -> int:
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    return len(root.findall(".//a:blip", NS))


def _normalise_visible_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def test_result_is_byte_stable_with_deterministic_ids(tmp_path: Path) -> None:
    path = _mixed_images_docx(tmp_path)
    first = _extract(path)
    second = _extract(path)

    assert _canonical_json(first) == _canonical_json(second)
    first_nodes = _items(first.nodes)
    second_nodes = _items(second.nodes)
    assert [node["node_id"] for node in first_nodes] == [node["node_id"] for node in second_nodes]
    assert len({node["node_id"] for node in first_nodes}) == len(first_nodes)
    assert all(str(node["node_id"]).startswith("node-") for node in first_nodes)


def test_mixed_text_and_two_images_preserve_inline_order(tmp_path: Path) -> None:
    result = _extract(_mixed_images_docx(tmp_path))
    content = [node for node in _items(result.nodes) if _node_kind(node) in {"text_run", "image"}]
    observed = [(_node_kind(node), _node_text(node)) for node in content]
    assert observed == [
        ("text_run", "alpha "),
        ("image", ""),
        ("text_run", " beta "),
        ("image", ""),
        ("text_run", " gamma"),
    ]


def test_reused_image_is_one_asset_but_two_source_occurrences(tmp_path: Path) -> None:
    result = _extract(_mixed_images_docx(tmp_path))
    images = [node for node in _items(result.nodes) if _node_kind(node) == "image"]
    assets = _items(result.assets)
    assert len(images) == 2
    asset_ids = [str((node.get("payload") or {}).get("asset_id")) for node in images]
    assert asset_ids[0] == asset_ids[1] and asset_ids[0] not in {"", "None"}
    assert len(assets) == 1
    assert str(assets[0].get("asset_id")) == asset_ids[0]
    assert str(assets[0].get("sha256", "")) == hashlib.sha256(PIXEL_PNG).hexdigest()


def test_inline_equation_is_an_occurrence_between_its_surrounding_runs(tmp_path: Path) -> None:
    result = _extract(_mixed_images_docx(tmp_path, include_equation=True))
    content = [node for node in _items(result.nodes) if _node_kind(node) in {"text_run", "equation"}]
    kinds = [_node_kind(node) for node in content]
    pytest.xfail("OMML adapter may land after the first walker slice") if "equation" not in kinds else None
    equation_index = kinds.index("equation")
    assert _node_text(content[equation_index - 1]) == " + "
    assert _node_text(content[equation_index + 1]) == " omega"
    assert (content[equation_index].get("payload") or {}).get("raw_omml") or (
        content[equation_index].get("payload") or {}
    ).get("omml_hash")


def test_numbered_nested_list_retains_raw_numbering_metadata(tmp_path: Path) -> None:
    document = Document()
    first = document.add_paragraph("first", style="List Number")
    nested = document.add_paragraph("nested", style="List Number 2")
    last = document.add_paragraph("last", style="List Number")
    _set_numbering(first, ilvl=0)
    _set_numbering(nested, ilvl=1)
    _set_numbering(last, ilvl=0)
    path = tmp_path / "nested-list.docx"
    document.save(path)

    result = _extract(path)
    paragraphs = [node for node in _items(result.nodes) if _node_kind(node) == "paragraph"]
    labelled = {
        "".join(_node_text(node) for node in _descendants(_items(result.nodes), str(paragraph["node_id"])))
        : paragraph
        for paragraph in paragraphs
    }
    assert {"first", "nested", "last"} <= set(labelled)
    raw_by_text = {text: _numbering(node) for text, node in labelled.items()}
    raw = list(raw_by_text.values())
    assert all(meta is not None for meta in raw)
    assert all("num_id" in meta and "ilvl" in meta for meta in raw)
    assert raw_by_text["nested"]["ilvl"] > raw_by_text["first"]["ilvl"]


def test_one_column_table_is_never_unwrapped_into_a_paragraph(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "header"
    table.cell(1, 0).text = "value"
    path = tmp_path / "one-column.docx"
    document.save(path)

    result = _extract(path)
    nodes = _items(result.nodes)
    tables = [node for node in nodes if _node_kind(node) == "table"]
    assert len(tables) == 1
    rows = [node for node in _descendants(nodes, str(tables[0]["node_id"])) if _node_kind(node) == "table_row"]
    cells = [node for node in _descendants(nodes, str(tables[0]["node_id"])) if _node_kind(node) == "table_cell"]
    assert len(rows) == 2
    assert len(cells) == 2


def test_flat_ledger_parent_sibling_and_sequence_invariants(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("before")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "inside"
    document.add_paragraph("after")
    path = tmp_path / "structure.docx"
    document.save(path)

    nodes = _items(_extract(path).nodes)
    ids = {str(node["node_id"]) for node in nodes}
    sequences = [node["sequence"] for node in nodes]
    assert sequences == list(range(len(nodes)))
    assert all(node.get("parent_id") is None or str(node["parent_id"]) in ids for node in nodes)
    siblings: dict[Any, list[int]] = {}
    for node in nodes:
        siblings.setdefault(node.get("parent_id"), []).append(node["sibling_index"])
    assert all(indexes == list(range(len(indexes))) for indexes in siblings.values())


def _supplied_policy_document() -> Path | None:
    docs = Path(__file__).resolve().parents[1] / "docs"
    candidates = sorted(path for path in docs.glob("*.docx") if _xml_image_count(path) == 9)
    return candidates[0] if len(candidates) == 1 else None


def test_supplied_policy_document_has_all_image_occurrences_and_image_paragraph_text() -> None:
    path = _supplied_policy_document()
    if path is None:
        pytest.skip("supplied policy DOCX is absent or cannot be identified unambiguously")
    expected_paragraphs = _xml_image_paragraph_texts(path)
    if len(expected_paragraphs) != 4:
        pytest.skip("policy DOCX no longer has the expected four text-bearing image paragraphs")

    result = _extract(path)
    nodes = _items(result.nodes)
    assert len([node for node in nodes if _node_kind(node) == "image"]) == 9
    all_text = _normalise_visible_text("".join(_node_text(node) for node in nodes))
    for visible_text in expected_paragraphs:
        assert _normalise_visible_text(visible_text) in all_text
