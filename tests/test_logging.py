"""Logging contract tests: explicit CLI setup and no document-content leakage."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from docx_pipeline.cli import _run
from docx_pipeline.logging import configure_logging, get_logger


def test_configure_logging_writes_requested_file(tmp_path: Path) -> None:
    log_file = tmp_path / "nested" / "pipeline.log"
    configure_logging("DEBUG", log_file)
    get_logger().debug("configured logging test")
    assert "configured logging test" in log_file.read_text(encoding="utf-8")


def test_get_logger_does_not_configure_process_sinks() -> None:
    """Binding a library logger is safe for host applications."""
    from loguru import logger

    before = len(logger._core.handlers)  # type: ignore[attr-defined]
    get_logger(operation="test")
    assert len(logger._core.handlers) == before  # type: ignore[attr-defined]


def test_cli_logging_flags_and_no_document_text_leakage(tmp_path: Path, capsys: object) -> None:
    secret = "DO_NOT_LOG_THIS_DOCUMENT_TEXT"
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph(secret)
    document.save(source)
    output = tmp_path / "result.json"
    log_file = tmp_path / "logs" / "pipeline.jsonl"

    assert _run(["extract", str(source), "--output", str(output), "--log-level", "DEBUG", "--log-file", str(log_file), "--log-json"]) == 0
    captured = capsys.readouterr()  # type: ignore[attr-defined]
    rendered = captured.err + log_file.read_text(encoding="utf-8")
    assert "Extraction completed" in rendered
    assert secret not in rendered


def test_validate_accepts_logging_flags(tmp_path: Path) -> None:
    document = Document()
    source = tmp_path / "source.docx"
    document.save(source)
    artifact = tmp_path / "result.json"
    assert _run(["extract", str(source), "--output", str(artifact), "--log-level", "ERROR"]) == 0
    assert _run(["validate", str(artifact), "--log-level", "ERROR", "--log-json"]) == 0
