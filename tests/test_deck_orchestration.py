"""Transaction, image scheduling, CLI, and independent PPTX checks."""

from __future__ import annotations

import asyncio
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from presentation_pipeline import DeckGenerationConfig, generate_deck
from presentation_pipeline import cli
from presentation_pipeline.deck import DeckPromotionError, _add_generated_images, _promote_staging_root
from presentation_pipeline.images import GeneratedImage
from presentation_pipeline.planning import PresentationRequirements
from presentation_pipeline.planning.models import SlidePurpose
from presentation_pipeline.rendering.verification import PptxVerificationReport, VerificationDiagnostic, verify_pptx
from presentation_pipeline.synthesis import PresentationContent, SlideContent

from tests.test_deck_generation import _Generator


class _ImageGenerator:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []
        self.in_flight = 0
        self.maximum_in_flight = 0

    async def generate_illustration(self, **kwargs):
        self.calls.append(kwargs)
        self.in_flight += 1
        self.maximum_in_flight = max(self.maximum_in_flight, self.in_flight)
        # Reverse delays make completion order differ from candidate order.
        await asyncio.sleep(0.01 * (4 - int(kwargs["slide_id"].split("-")[-1])))
        self.in_flight -= 1
        return GeneratedImage(
            slide_id=kwargs["slide_id"],
            path=Path(kwargs["output_dir"]) / f"{kwargs['slide_id']}.png",
            prompt=f"prompt {kwargs['slide_id']}",
            model="test-image-model",
        )


def _image_context(number: int):
    slide_id = f"slide-{number}"
    slide = SimpleNamespace(
        slide_id=slide_id,
        purpose=SlidePurpose.CONTENT,
        preferred_content_forms=["image"],
        title=f"Title {number}",
        message=f"Message {number}",
    )
    return SimpleNamespace(slide=slide)


def test_generated_image_selection_is_capped_before_calls_and_ordered(tmp_path) -> None:
    contexts = [_image_context(index) for index in range(1, 5)]
    content = PresentationContent(slides=[SlideContent(slide_id=context.slide.slide_id, elements=[]) for context in contexts])
    resolved = {context.slide.slide_id: [] for context in contexts}
    generator = _ImageGenerator()

    outcomes = asyncio.run(_add_generated_images(
        contexts, content, resolved, generator, tmp_path, DeckGenerationConfig(generate_images=True, max_generated_images=3, image_concurrency=2), [], audience="Leaders", tone="direct"
    ))

    assert [call["slide_id"] for call in generator.calls] == ["slide-1", "slide-2", "slide-3"]
    assert generator.maximum_in_flight <= 2
    assert [outcome["slide_id"] for outcome in outcomes] == ["slide-1", "slide-2", "slide-3"]
    assert all(call["audience"] == "Leaders" and call["tone"] == "direct" for call in generator.calls)
    assert all("prompt_hash" in outcome and "prompt" not in outcome for outcome in outcomes)


def _source_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Quarterly update", level=1)
    document.add_paragraph("Revenue increased during the quarter.")
    document.save(path)
    return path


def _generate(source: Path, output: Path, *, overwrite: bool = False):
    return asyncio.run(generate_deck(
        [source], PresentationRequirements(goal="Brief leadership", audience="Leadership", target_slide_count=2), _Generator(),
        output_dir=output, config=DeckGenerationConfig(), extraction_workers=1, llm_concurrency=1, overwrite=overwrite,
    ))


def test_output_transaction_rejects_existing_paths_and_keeps_json_public(tmp_path) -> None:
    source = _source_docx(tmp_path / "briefing.docx")
    file_target = tmp_path / "output-file"
    file_target.write_text("not a directory", encoding="utf-8")
    with pytest.raises(ValueError, match="directory"):
        _generate(source, file_target)

    directory_target = tmp_path / "output-directory"
    directory_target.mkdir()
    with pytest.raises(FileExistsError, match="already exists"):
        _generate(source, directory_target)

    symlink_target = tmp_path / "output-symlink"
    try:
        symlink_target.symlink_to(directory_target, target_is_directory=True)
    except OSError:  # pragma: no cover - platforms without symlink permission
        pytest.skip("symlinks are unavailable")
    with pytest.raises(ValueError, match="symlink"):
        _generate(source, symlink_target)

    result = _generate(source, tmp_path / "published")
    for path in (result.layout_path, result.report_path):
        assert ".staging-" not in path.read_text(encoding="utf-8")


def test_failed_promotion_rolls_back_preexisting_output(tmp_path, monkeypatch) -> None:
    source = _source_docx(tmp_path / "briefing.docx")
    output = tmp_path / "published"
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("previous result", encoding="utf-8")

    def reject(*_args, **kwargs):
        return PptxVerificationReport(
            output_path=str(kwargs["reported_path"]), verified=False, slide_count=None, used_soffice=False,
            diagnostics=(VerificationDiagnostic("error", "TEST_REJECT", "Test verification rejection."),),
        )

    monkeypatch.setattr("presentation_pipeline.rendering.verification.verify_pptx", reject)
    with pytest.raises(RuntimeError, match="verification failed"):
        _generate(source, output, overwrite=True)
    assert sentinel.read_text(encoding="utf-8") == "previous result"
    assert not list(tmp_path.glob(".published.staging-*"))


def test_mid_promotion_failure_restores_prior_directory(tmp_path, monkeypatch) -> None:
    output = tmp_path / "published"
    output.mkdir()
    (output / "keep.txt").write_text("previous result", encoding="utf-8")
    staging = tmp_path / ".published.staging-test"
    staging.mkdir()
    (staging / "deck.pptx").write_bytes(b"new result")
    original_replace = Path.replace

    def fail_staging_promotion(path: Path, target: Path):
        if path == staging:
            raise OSError("injected promotion failure")
        return original_replace(path, target)

    monkeypatch.setattr(Path, "replace", fail_staging_promotion)
    with pytest.raises(DeckPromotionError, match="could not promote"):
        _promote_staging_root(staging, output, overwrite=True)
    assert (output / "keep.txt").read_text(encoding="utf-8") == "previous result"
    assert staging.is_dir()
    assert not list(tmp_path.glob(".published.backup-*"))


def test_independent_verifier_checks_ooxml_even_without_soffice(tmp_path) -> None:
    from pptx import Presentation

    valid = tmp_path / "valid.pptx"
    Presentation().save(valid)
    report = verify_pptx(valid, expected_slide_count=0, try_soffice=False)
    assert report.verified and report.slide_count == 0 and not report.used_soffice
    assert report.verifier_mode == "ooxml_zip"

    invalid = tmp_path / "invalid.pptx"
    invalid.write_text("not a zip", encoding="utf-8")
    report = verify_pptx(invalid, try_soffice=False)
    assert not report.verified
    assert report.diagnostics[0].code == "PPTX_NOT_ZIP"


def test_independent_verifier_rejects_missing_and_malformed_required_parts(tmp_path) -> None:
    from pptx import Presentation

    valid = tmp_path / "valid.pptx"
    Presentation().save(valid)
    missing = tmp_path / "missing-root-rels.pptx"
    malformed = tmp_path / "malformed-presentation.pptx"
    with ZipFile(valid) as source, ZipFile(missing, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename != "_rels/.rels":
                target.writestr(item, source.read(item.filename))
    with ZipFile(valid) as source, ZipFile(malformed, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = b"<broken" if item.filename == "ppt/presentation.xml" else source.read(item.filename)
            target.writestr(item, payload)

    missing_report = verify_pptx(missing, try_soffice=False)
    malformed_report = verify_pptx(malformed, try_soffice=False)
    assert not missing_report.verified
    assert any(item.code == "PPTX_REQUIRED_PART_MISSING" for item in missing_report.diagnostics)
    assert not malformed_report.verified
    assert any(item.code == "PPTX_OOXML_INVALID" for item in malformed_report.diagnostics)


def test_cli_exposes_and_forwards_image_concurrency(monkeypatch, tmp_path) -> None:
    parsed = cli._parser().parse_args([
        "generate", "source.docx", "--goal", "Goal", "--audience", "Audience", "--slides", "2", "--output-dir", str(tmp_path),
        "--generate-images", "--image-concurrency", "3",
    ])
    received: dict[str, object] = {}

    async def fake_generate_deck(*_args, **kwargs):
        received.update(kwargs)
        return SimpleNamespace(pptx_path=tmp_path / "deck.pptx")

    monkeypatch.setattr(cli, "generate_deck", fake_generate_deck)
    monkeypatch.setattr(cli, "OpenAIStructuredGenerator", lambda **_kwargs: object())
    monkeypatch.setattr(cli, "OpenAIImageGenerator", lambda **_kwargs: object())
    assert asyncio.run(cli._run(parsed)) == tmp_path / "deck.pptx"
    assert received["config"].image_concurrency == 3
