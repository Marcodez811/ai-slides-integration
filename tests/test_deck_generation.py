"""Offline DOCX-to-PPTX acceptance coverage."""

from __future__ import annotations

import asyncio
import json

from docx import Document
from pptx import Presentation

from presentation_pipeline import DeckGenerationConfig, generate_deck
from presentation_pipeline.planning import PresentationRequirements
from presentation_pipeline.planning.models import EvidenceSelection, PresentationOutline
from presentation_pipeline.retrieval.models import LocalCandidateSelection
from presentation_pipeline.synthesis import SlideContent
from presentation_pipeline.understanding.models import ChunkDigest, DigestFragment, DocumentDigest


class _Generator:
    """Return valid, source-backed output without a network provider."""

    async def generate(self, *, input_data, response_model, **_kwargs):
        if response_model is DigestFragment:
            document = input_data["document"]
            fragment = input_data["fragments"][0]
            return {
                "doc_id": document["doc_id"],
                "fragment_id": "offline-reduction",
                "summary": fragment["summary"],
                "topics": fragment["topics"],
                "key_facts": fragment["key_facts"],
            }
        if response_model in (DocumentDigest, ChunkDigest):
            document = input_data["document"]
            if response_model is DocumentDigest and "fragments" in input_data:
                fragment = input_data["fragments"][0]
                return {
                    "doc_id": document["doc_id"],
                    "summary": fragment["summary"],
                    "topics": fragment["topics"],
                    "key_facts": fragment["key_facts"],
                }
            evidence = input_data["evidence"][0]
            output = {
                "doc_id": document["doc_id"],
                "summary": "Briefing summary",
                "topics": [{
                    "topic": "Briefing",
                    "summary": "Quarterly source update",
                    "evidence": [{"doc_id": document["doc_id"], "evidence_ids": [evidence["evidence_id"]]}],
                }],
                "key_facts": [{
                    "claim": evidence.get("text", "Source fact"),
                    "evidence": [{"doc_id": document["doc_id"], "evidence_ids": [evidence["evidence_id"]]}],
                }],
            }
            if response_model is ChunkDigest:
                output["window_id"] = input_data["window"]["window_id"]
            return output
        if response_model is LocalCandidateSelection:
            if "candidate_evidence" in input_data:
                evidence = input_data["candidate_evidence"][0]
                return {"candidates": [{
                    "doc_id": evidence["doc_id"],
                    "evidence_id": evidence["evidence_id"],
                    "reason": "Reduced source briefing evidence",
                }]}
            evidence = input_data["evidence"][0]
            return {"candidates": [{
                "doc_id": input_data["window"]["doc_id"],
                "evidence_id": evidence["evidence_id"],
                "reason": "Source briefing evidence",
            }]}
        if response_model is EvidenceSelection:
            evidence = input_data["candidate_evidence"][0]
            return {"selected": [{
                "doc_id": evidence["doc_id"],
                "evidence_id": evidence["evidence_id"],
                "reason": "Source briefing evidence",
            }]}
        if response_model is PresentationOutline:
            selected = input_data["selected_evidence"][0]
            target = input_data["requirements"]["target_slide_count"]
            slides = [{"slide_id": "slide-1", "title": "Quarterly update", "purpose": "title", "message": "Quarterly update"}]
            slides.extend({
                "slide_id": f"slide-{number}",
                "title": f"Evidence {number - 1}",
                "purpose": "content",
                "message": "Revenue increased during the quarter.",
                "preferred_content_forms": ["text"],
                "evidence": [{"doc_id": selected["doc_id"], "evidence_ids": [selected["evidence_id"]]}],
            } for number in range(2, target + 1))
            return {
                "title": "Quarterly update",
                "objective": "Brief leadership",
                "narrative": "Start with the update.",
                "sections": [{
                    "section_id": "section-1",
                    "title": "Update",
                    "purpose": "Brief leadership",
                    "slides": slides,
                }],
            }
        if response_model is SlideContent:
            slide = input_data["slide"]
            if slide["purpose"] == "title":
                return {"slide_id": slide["slide_id"], "elements": []}
            evidence = input_data["evidence"][0]
            return {"slide_id": slide["slide_id"], "elements": [{
                "kind": "text",
                "text": "Revenue increased during the quarter.",
                "evidence": [{"doc_id": evidence["doc_id"], "evidence_ids": [evidence["evidence_id"]]}],
            }]}
        raise AssertionError(f"unexpected response model: {response_model}")


def test_real_docx_reaches_a_reopenable_pptx_with_audit_artifacts(tmp_path) -> None:
    source = tmp_path / "briefing.docx"
    document = Document()
    document.add_heading("Quarterly update", level=1)
    document.add_paragraph("Revenue increased during the quarter.")
    document.save(source)

    result = asyncio.run(generate_deck(
        [source],
        PresentationRequirements(goal="Brief leadership", audience="Leadership", target_slide_count=2),
        _Generator(),
        output_dir=tmp_path / "output",
        config=DeckGenerationConfig(),
        extraction_workers=1,
        llm_concurrency=1,
    ))

    presentation = Presentation(result.pptx_path)
    assert len(presentation.slides) == 2
    assert result.outline_path.is_file()
    assert result.content_path.is_file()
    assert result.layout_path.is_file()
    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["semantic_slide_count"] == 2
    assert report["physical_slide_count"] == 2
